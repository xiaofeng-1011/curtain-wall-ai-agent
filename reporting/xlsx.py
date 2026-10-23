"""Excel 比选表：面向同事复核的主要交付物（技术追溯包 ZIP 另行提供）。"""
from __future__ import annotations

import io

from .report import SCOPE_STATEMENT, banner, comparison_table, param_rows


def build_xlsx(run: dict, project_name: str | None = None, *, current: bool = True) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "比选表"
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EEF2F6")
    ws.append(["已覆盖铝型材理论材料金额比选表" + ("" if run["execution_state"] == "COMPLETE" else "（部分结果）")])
    ws["A1"].font = Font(bold=True, size=14)
    info = [("项目", project_name or run.get("project_name") or ""), ("数据状态", banner(run)),
            ("计算时间（UTC）", run.get("started_at")), ("运行编号", run.get("run_id")),
            ("结果对应版本", "当前参数" if current else "历史版本（之后参数已修改）"),
            ("完成情况", run.get("completeness_text")), ("计量口径", SCOPE_STATEMENT)]
    for k, v in info:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    rows = comparison_table(run)
    cols = ["指标", "单位"] + [c for c in ("原方案", "调整方案", "差值") if any(r.get(c) is not None for r in rows)]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(ws.max_row, c).font = bold
        ws.cell(ws.max_row, c).fill = head_fill
    for r in rows:
        ws.append([r.get(c) for c in cols])
    pct = ((run.get("comparison") or {}).get("percent_metrics") or {}).get("percent_total_cost")
    if pct:
        ws.append(["总金额变化率", "%", None, None,
                   pct["display_value"] if pct.get("value") is not None else "不可计算（原方案金额为零）"])
    comp = run.get("comparison") or {}
    if comp.get("changed_fields"):
        ws.append([])
        ws.append(["变更项", "原方案", "调整方案"])
        for c in range(1, 4):
            ws.cell(ws.max_row, c).font = bold
        for c in comp["changed_fields"]:
            ws.append([c["label"], c["baseline"], c["candidate"]])
    ws.append([])
    ws.append(["待确认事项"])
    ws.cell(ws.max_row, 1).font = bold
    for it in run.get("review_items") or []:
        ws.append(["□", it])
    ws.append([])
    ws.append(["说明", "差值 = 调整方案 − 原方案；负号表示减少，不等于工程优化。结果不含玻璃、五金、加工、损耗、运输、安装与税费，"
                     "不作结构安全验算或工程签审；型材规格适用性需工程师确认。"])
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 22
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws2 = wb.create_sheet("输入参数与来源")
    prs = param_rows(run)
    if prs:
        keys = list(prs[0].keys())
        ws2.append(keys)
        for c in range(1, len(keys) + 1):
            ws2.cell(1, c).font = bold
        for r in prs:
            ws2.append([r.get(k) for k in keys])
    ws2.append([])
    for role, title in (("baseline", "原方案"), ("candidate", "调整方案")):
        s = (run.get("snapshots") or {}).get(role)
        if not s:
            continue
        for comp_key, label in (("column", "立柱"), ("beam", "横梁")):
            p = (s.get("resolved_profiles") or {}).get(comp_key) or {}
            pr = (s.get("resolved_prices") or {}).get(comp_key) or {}
            ws2.append([f"{title}{label}", p.get("profile_name"), f"线重 {p.get('linear_mass') or '缺失'} kg/m",
                        f"单价 {pr.get('unit_price') if pr.get('unit_price') is not None else '缺失'} 元/kg",
                        pr.get("tax_basis")])
    ws2.append([])
    ws2.append(["来源编号", "来源状态", "来源类型", "审核状态", "出处"])
    provs = {}
    for s in (run.get("snapshots") or {}).values():
        provs.update(s.get("provenance") or {})
    for pid, p in sorted(provs.items()):
        ws2.append([pid, p.get("source_status"), p.get("source_type"), p.get("review_status"), p.get("source_ref")])

    ws3 = wb.create_sheet("计算过程")
    ws3.append(["方案", "公式", "代入", "结果", "单位"])
    for c in range(1, 6):
        ws3.cell(1, c).font = bold
    for role, title in (("baseline", "原方案"), ("candidate", "调整方案")):
        res = (run.get("results") or {}).get(role) or {}
        for mod in ("geometry", "length", "mass", "cost"):
            for st in (res.get(mod) or {}).get("steps", []):
                ws3.append([title, st["formula_id"], st["expression"], st["exact_value"], st["unit"]])
    for st in (comp.get("steps") or []):
        ws3.append(["差值", st["formula_id"], st["expression"], st["exact_value"], st["unit"]])
    ws3.column_dimensions["C"].width = 48
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
