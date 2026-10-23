"""改版方案 §9.4：规则、追溯和来源（T31–T38）。"""
import io

import pytest

from agent import documents
from agent import projects as P
from tests.chat_helpers import Chat, disp


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv("MW_FORCE_FALLBACK", "1")


def _rules(c):
    pid = c.project()["project_id"]
    c.a.add_rule(pid, title="分格宽演示条件", lhs="grid_width_m", operator="<=", threshold="1.5", unit="m")
    c.a.add_rule(pid, title="板块面积演示条件", lhs="panel_area_m2", operator="<=", threshold="2", unit="m²")
    c.a.add_rule(pid, title="层高演示条件", lhs="floor_height_m", operator="<=", threshold="3.6", unit="m")


def test_t31_t32_t23_demo_rules_all_executed(tmp_path):
    c = Chat(tmp_path).build_t10()
    _rules(c)
    c.say("把分格宽改为2m，与原方案比较")
    r = c.run()
    cand = {x["title"]: x["status"] for x in r["rule_results"]["candidate"]["rule_results"]}
    assert cand == {"分格宽演示条件": "DEMO_ATTENTION", "板块面积演示条件": "DEMO_ATTENTION",
                    "层高演示条件": "NEED_PARAMETERS"}  # 未提供层高 → 待补参数，不按满足处理
    base = {x["title"]: x["status"] for x in r["rule_results"]["baseline"]["rule_results"]}
    assert base["分格宽演示条件"] == "DEMO_MET" and base["板块面积演示条件"] == "DEMO_MET"
    assert r["coverage"]["formal_clause_count"] == 0 and r["coverage"]["required_count"] == 3
    assert disp(r, "candidate", "total_cost_cny") == "7,326.00"  # 规则缺参不阻断成本


def test_t32_zero_rules_is_not_pass(tmp_path):
    c = Chat(tmp_path).build_t10()
    r = c.run()
    assert "rule_check" not in r["module_statuses"]
    assert "未执行正式规范核查" in r["report"]["banner"]
    assert "全部合规" not in r["report"]["markdown"] and "全部通过" not in r["report"]["markdown"]


def test_t34_rules_without_price(tmp_path):
    c = Chat(tmp_path).new()
    c.say("宽18m、高6m，分格宽1.2m、高1.5m")
    _rules(c)
    c.say("执行规则核查")
    c.click("确认并计算")
    r = c.run()
    st = {x["title"]: x["status"] for x in r["rule_results"]["baseline"]["rule_results"]}
    assert st["分格宽演示条件"] == "DEMO_MET" and st["板块面积演示条件"] == "DEMO_MET"


def test_t36_sources_per_field(tmp_path):
    c = Chat(tmp_path).new()
    c.say("宽18m、高6m，分格宽1.2m、高1.5m，这是测试数据")
    c.say("立柱2.7kg/m 25元/kg，横梁1.3kg/m 28元/kg，来自真实项目报价")
    p = c.project()
    gprov = p["provenance"][p["schemes"]["baseline"]["params"]["width_m"]["provenance_ids"][0]]
    mprov = p["provenance"][P.material_of(p, "baseline", "column")["provenance_id"]]
    assert gprov["source_status"] == "synthetic"
    assert mprov["source_status"] == "pending" and "真实项目" in mprov["notes"]  # 不自行升级为 verified
    for pv in p["provenance"].values():
        assert pv["source_status"] != "verified"
        assert {"source_status", "source_type", "review_status", "generation_method"} <= set(pv)
    c.say("算金额")
    c.click("确认并计算")
    r = c.run()
    assert r["contains_synthetic"] and r["contains_pending"]
    m = r["results"]["baseline"]["cost"]["metrics"]["total_cost_cny"]
    assert set(m["source_statuses"]) == {"synthetic", "pending"} and m["review_status"] == "pending_professional_review"


def test_t37_traceability(tmp_path):
    c = Chat(tmp_path).build_t10()
    r = c.run()
    s = r["snapshots"]["baseline"]
    assert s["project_id"] == c.project()["project_id"] and s["scheme_rev"] >= 1
    cc = r["results"]["baseline"]["cost"]["metrics"]["column_cost_cny"]
    assert any("resolved_prices.column.unit_price" in x for x in cc["input_refs"])
    assert r["tool_versions"]["TL10"] and r["started_at"]
    pid = s["resolved_profiles"]["column"]["provenance_ids"][0]
    assert s["provenance"][pid]["source_ref"].startswith("你在本轮输入")


def test_t38_xlsx_export(tmp_path):
    import openpyxl
    from reporting.xlsx import build_xlsx
    c = Chat(tmp_path).build_t10()
    c.say("把分格宽改为1.5m，与原方案比较")
    wb = openpyxl.load_workbook(io.BytesIO(build_xlsx(c.run(), "手工测试甲")))
    assert wb.sheetnames == ["比选表", "输入参数与来源", "计算过程"]
    cells = [str(v) for row in wb["比选表"].iter_rows(values_only=True) for v in row if v is not None]
    joined = " ".join(cells)
    assert "原方案" in joined and "调整方案" in joined and "9,756.00" in joined and "8,541.00" in joined
    assert "演示数据" in joined and "不作结构安全验算" in joined and "未税" in joined or "仅立柱横梁" in joined
    c.say("导出比选表")
    assert c.last()["kind"] == "export"


def test_t27_query_uploaded_docs(tmp_path):
    c = Chat(tmp_path).new()
    c.say("玻璃幕墙立柱截面厚度有什么要求？")
    text = ("5.6.2 立柱截面的主要受力部分的厚度应符合下列规定：铝合金型材截面开口部位的厚度不应小于 3.0mm，"
            "闭口部位的厚度不应小于 2.5mm。" * 3)
    c.say("", files=[("测试资料.txt", text.encode("utf-8"))])
    assert documents and c.ws.counts()["documents"] == 1
    c.click("用刚才的问题检索")
    m = c.last("query")
    run = c.ws.get_run(m["card"]["run_id"])
    hit = run["retrieval"]["candidates"][0]
    assert hit["source_status"] == "pending" and hit["usable_for_determination"] is False
    assert "第 1 页" in hit["location"] and "未经审核" in m["text"]
    assert c.ws.counts()["projects"] == 0
