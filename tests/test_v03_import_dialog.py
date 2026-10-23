"""改版方案 §9.3：导入、对话、检索与边界（T22–T30）。"""
import io
import json
import sys
import types

import pytest

from agent import extraction
from agent import projects as P
from tests.chat_helpers import Chat, disp


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    monkeypatch.setenv("MW_FORCE_FALLBACK", "1")


def _xlsx(rows, sheet="参数"):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import_msg(c):
    return next(m for m in reversed(c.conv["messages"]) if m["kind"] == "import")


def test_t22_xlsx_wide_layout(tmp_path):
    data = _xlsx([["方案", "宽度(m)", "总高(m)", "分格宽(mm)", "分格高(m)", "立柱线重(kg/m)", "横梁线重(kg/m)",
                   "立柱单价(元/kg)", "横梁单价(元/kg)", "备注"],
                  ["原方案", 18, 6, 1200, 1.5, 2.7, 1.3, 25, 28, "测试"]])
    c = Chat(tmp_path).new()
    c.say("", files=[("参数表.xlsx", data)])
    m = _import_msg(c)
    pv = c.conv["imports"][m["card"]["import_id"]]
    assert pv["layout"] == "wide" and not pv["missing"]
    assert any("备注" in n for n in pv["notes"])
    assert c.ws.counts()["projects"] == 0  # 预览不写数据
    c.conv = c.a.apply_import(c.conv["conversation_id"], m["id"], pv["import_id"], mapping={}, mode="new",
                              declared="synthetic")
    p = c.project()
    assert P.user_value(p, "baseline", "grid_width_m")[0] == "1.2"
    prov = p["provenance"][p["schemes"]["baseline"]["params"]["width_m"]["provenance_ids"][0]]
    assert "参数表.xlsx｜参数｜第 2 行第 2 列" in prov["source_ref"] and prov["source_status"] == "synthetic"
    assert (c.ws.root / "attachments" / p["project_id"] / "参数表.xlsx").exists()
    c.click("计算材料金额")
    c.click("确认并计算")
    assert disp(c.run(), "baseline", "total_cost_cny") == "9,756.00"


def test_t23_missing_price_import(tmp_path):
    data = _xlsx([["参数名称", "数值", "单位"], ["幕墙总宽", 18, "m"], ["幕墙总高", 6, "m"], ["分格宽", 1.2, "m"],
                  ["分格高", 1.5, "m"], ["立柱线重", 2.7, "kg/m"], ["横梁线重", 1.3, "kg/m"]])
    c = Chat(tmp_path).new()
    c.say("", files=[("缺价.xlsx", data)])
    m = _import_msg(c)
    pv = c.conv["imports"][m["card"]["import_id"]]
    assert pv["layout"] == "params" and "立柱单价" in pv["missing"]
    c.conv = c.a.apply_import(c.conv["conversation_id"], m["id"], pv["import_id"], mapping={}, mode="new",
                              declared=None)
    c.click("计算型材用量")
    c.click("确认并计算")
    assert disp(c.run(), "baseline", "total_mass_kg") == "376.200"


def test_t24_ambiguous_column_needs_choice(tmp_path):
    c = Chat(tmp_path).new()
    c.say("幕墙总宽(m)\t总高(m)\t尺寸\n18\t6\t1.2")
    m = _import_msg(c)
    pv = c.conv["imports"][m["card"]["import_id"]]
    assert [a["column"] for a in pv["ambiguous"]] == ["尺寸"]
    c.conv = c.a.apply_import(c.conv["conversation_id"], m["id"], pv["import_id"], mapping={}, mode="new",
                              declared=None)
    p = c.project()
    assert P.user_value(p, "baseline", "grid_width_m")[0] is None  # 未选择前不写入
    assert P.user_value(p, "baseline", "width_m")[0] == "18"


def test_t24_ambiguous_column_mapped(tmp_path):
    c = Chat(tmp_path).new()
    c.say("幕墙总宽\t总高\t尺寸\n18\t6\t1.2m")
    m = _import_msg(c)
    pv = c.conv["imports"][m["card"]["import_id"]]
    key = pv["ambiguous"][0]["key"]
    c.conv = c.a.apply_import(c.conv["conversation_id"], m["id"], pv["import_id"], mapping={key: "grid_width_m"},
                              mode="new", declared=None)
    assert P.user_value(c.project(), "baseline", "grid_width_m")[0] == "1.2"


def test_materials_layout(tmp_path):
    c = Chat(tmp_path).new()
    c.say("宽18m、高6m，分格宽1.2m、高1.5m")
    c.say("用途\t名称\t线重(kg/m)\t单价\t价格单位\t税费说明\n立柱\t立柱甲\t2.7\t25000\t元/吨\t未税\n横梁\t横梁乙\t1.3\t28\t元/kg\t未税")
    m = _import_msg(c)
    pv = c.conv["imports"][m["card"]["import_id"]]
    assert pv["layout"] == "materials"
    c.conv = c.a.apply_import(c.conv["conversation_id"], m["id"], pv["import_id"], mapping={}, mode="current",
                              declared=None)
    p = c.project()
    assert P.material_of(p, "baseline", "column")["profile_name"] == "立柱甲"
    assert P.user_value(p, "baseline", "column_unit_price")[0] == "25"


def test_template_has_no_values():
    import openpyxl
    from agent.importer import template_xlsx
    wb = openpyxl.load_workbook(io.BytesIO(template_xlsx()))
    for row in wb["方案参数"].iter_rows(min_row=2, values_only=True):
        assert row[1] is None
    for row in wb["材料报价"].iter_rows(min_row=2, values_only=True):
        assert row[2] is None and row[3] is None


def test_t25_chinese_numbers_and_relative(tmp_path):
    c = Chat(tmp_path).new()
    c.say("宽十八米、高六米，分格宽一米二、高一米五")
    p = c.project()
    assert [P.user_value(p, "baseline", f)[0] for f in ("width_m", "height_m", "grid_width_m", "grid_height_m")] == \
        ["18", "6", "1.2", "1.5"]
    c.say("分格高再增加 100 毫米")
    m = c.last("change")
    assert "分格高 1.5 → 1.6 m（1.5 m + 0.1 m）" in m["text"]


def test_t26_explain_uses_steps(tmp_path):
    c = Chat(tmp_path).build_t10()
    c.say("把分格宽改为1.5m，与原方案比较")
    c.say("为什么金额减少？")
    m = c.last("explain")
    assert "Lc = Nc × H = 16 × 6 = 96 m" in m["text"] and "Lc = Nc × H = 13 × 6 = 78 m" in m["text"]
    assert "结构安全" not in m["text"]


@pytest.mark.parametrize("text", ["单元式幕墙怎么计算", "完整结构安全么", "帮我读取 CAD 图纸"])
def test_t28_out_of_scope(tmp_path, text):
    c = Chat(tmp_path).build_t10()
    runs_before = c.ws.counts()["runs"]
    c.say(text)
    m = c.last()
    assert m["kind"] == "scope" and not any(ch.isdigit() for ch in m["text"].split("可以继续")[0])
    assert c.ws.counts()["runs"] == runs_before


def test_t29_old_result_marked(tmp_path):
    from agent.engine import is_current
    c = Chat(tmp_path).build_t10()
    first = c.run()
    c.say("把分格宽改为1.5m")
    assert not is_current(c.project(), first)


def test_t30_model_failure_keeps_input(tmp_path, monkeypatch):
    monkeypatch.delenv("MW_FORCE_FALLBACK", raising=False)
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test")
    monkeypatch.setenv("MW_PROVIDER", "deepseek")

    class Boom:
        def __init__(self, **kw):
            self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=self.create))

        def create(self, **kw):
            raise TimeoutError("timed out")

    monkeypatch.setitem(sys.modules, "openai", types.SimpleNamespace(OpenAI=Boom))
    c = Chat(tmp_path).new()
    c.say("宽18m、高6m，分格宽1.2m、高1.5m")
    assert P.user_value(c.project(), "baseline", "width_m")[0] == "18"  # 本地解析继续
    m = c.last("change")
    assert "在线理解暂不可用" in (m.get("note") or "")
    assert c.ws.get_setting("model_status")["ok"] is False
    assert c.conv["messages"][0]["text"] == "宽18m、高6m，分格宽1.2m、高1.5m"


def test_t30_invalid_model_json(tmp_path, monkeypatch):
    monkeypatch.delenv("MW_FORCE_FALLBACK", raising=False)
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test")
    monkeypatch.setenv("MW_PROVIDER", "deepseek")

    class Bad:
        def __init__(self, **kw):
            msg = types.SimpleNamespace(content=json.dumps({"intents": ["approve_safety"]}))
            out = types.SimpleNamespace(choices=[types.SimpleNamespace(message=msg, finish_reason="stop")])
            self.chat = types.SimpleNamespace(completions=types.SimpleNamespace(create=lambda **k: out))

    monkeypatch.setitem(sys.modules, "openai", types.SimpleNamespace(OpenAI=Bad))
    c = Chat(tmp_path).new()
    c.say("宽18m、高6m")
    assert P.user_value(c.project(), "baseline", "height_m")[0] == "6"
