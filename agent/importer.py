"""普通表格导入：粘贴文本、CSV、XLSX → 映射预览 → 用户确认后写入当前项目。

支持三种常见布局：
- 参数清单：参数名称 | 数值 | 单位 [| 方案名称]
- 一行一个方案：方案 | 幕墙宽(m) | 总高(m) | 分格宽 | …
- 材料报价：用途（立柱/横梁）| 名称 | 线重(kg/m) | 单价 | 价格单位 [| 出处 | 日期 | 税费说明]
无法判断含义的列（如“尺寸”）不会写入，只在预览中请用户选择；不要求 JSON 数据包或特殊编号。
"""
from __future__ import annotations

import csv
import io
import re

from domain.contracts import Code, Issue, new_id
from domain.fields import USER_FIELDS
from domain.units import parse_quantity

SYNONYMS = {
    "width_m": ["幕墙宽度", "幕墙宽", "立面宽度", "立面宽", "总宽度", "总宽", "宽度", "宽", "w"],
    "height_m": ["幕墙高度", "幕墙高", "立面高度", "立面高", "总高度", "总高", "高度", "高", "h"],
    "grid_width_m": ["分格宽度", "分格宽", "板块宽度", "板块宽", "b"],
    "grid_height_m": ["分格高度", "分格高", "板块高度", "板块高"],
    "floor_height_m": ["层高"],
    "reported_area_m2": ["申报面积", "立面面积", "面积"],
    "column_linear_mass": ["立柱线重", "立柱线密度", "立柱每米重", "立柱米重"],
    "beam_linear_mass": ["横梁线重", "横梁线密度", "横梁每米重", "横梁米重"],
    "column_unit_price": ["立柱单价", "立柱价格", "立柱报价"],
    "beam_unit_price": ["横梁单价", "横梁价格", "横梁报价"],
    "column_name": ["立柱名称", "立柱型号", "立柱型材"],
    "beam_name": ["横梁名称", "横梁型号", "横梁型材"],
}
AMBIGUOUS = ["尺寸", "规格", "大小", "长度", "数值2"]
UNIT_IN_HEADER = re.compile(r"^(.*?)\s*[（(\[／/]?\s*(mm|毫米|cm|厘米|m²|m2|平方米|kg/m|千克/米|公斤/米|元/kg|元/千克|元/公斤|元/吨|元/t|m|米)\s*[)）\]]?\s*$",
                            re.IGNORECASE)
PARAM_NAME_COLS = ["参数名称", "参数", "名称", "项目", "字段"]
VALUE_COLS = ["数值", "值", "取值", "数量"]
UNIT_COLS = ["单位"]
SCHEME_COLS = ["方案名称", "方案"]
USE_COLS = ["用途", "构件", "类型", "部位"]
MAT_COLS = {"name": ["名称", "型材名称", "型号"], "linear_mass": ["线重", "线密度", "每米重", "米重"],
            "unit_price": ["单价", "价格", "报价"], "price_unit": ["价格单位", "单价单位", "计价单位"],
            "source": ["出处", "来源"], "date": ["日期", "报价日期"], "tax": ["税费说明", "税率", "含税", "税"]}


def _clean(h) -> str:
    return re.sub(r"\s+", "", str(h or "")).strip()


def _split_header(h: str) -> tuple[str, str | None]:
    h = _clean(h)
    m = UNIT_IN_HEADER.match(h)
    if m and m.group(1):
        return m.group(1).rstrip("/／"), m.group(2)
    return h, None


def match_field(name: str) -> str | None:
    n = name.lower()
    variants = [n, re.sub(r"^(幕墙|立面)", "", n), re.sub(r"度$", "", n), re.sub(r"度$", "", re.sub(r"^(幕墙|立面)", "", n))]
    for v in variants:
        for f, syns in SYNONYMS.items():
            if v and v in [s.lower() for s in syns]:
                return f
    return None


def _cell(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        v = repr(v)
        if v.endswith(".0"):
            v = v[:-2]
    s = str(v).strip()
    return s or None


# ------------------------------------------------------------------ 读取
def read_paste(text: str) -> list[list[str]] | None:
    lines = [ln for ln in text.strip().splitlines() if ln.strip()]
    if len(lines) < 2:
        return None
    for delim in ("\t", "|", ",", "，"):
        rows = [[c.strip() for c in ln.strip().strip("|").split(delim)] for ln in lines]
        if all(len(r) >= 2 for r in rows) and len({len(r) for r in rows}) <= 2:
            return rows
    rows = [re.split(r"\s{2,}", ln.strip()) for ln in lines]
    if all(len(r) >= 2 for r in rows):
        return rows
    return None


def looks_like_table(text: str) -> bool:
    rows = read_paste(text)
    if not rows:
        return False
    header = [_split_header(h)[0] for h in rows[0]]
    known = sum(1 for h in header if match_field(h) or h in PARAM_NAME_COLS + VALUE_COLS + UNIT_COLS + USE_COLS
                + sum(MAT_COLS.values(), []) + AMBIGUOUS + SCHEME_COLS)
    return known >= 2 or (known >= 1 and "\t" in text)


def read_file(filename: str, data: bytes) -> dict:
    """返回 {sheets: {name: rows}}；读取失败返回 {error}。"""
    low = filename.lower()
    try:
        if low.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            return {"sheets": {"CSV": [row for row in csv.reader(io.StringIO(text))]}}
        if low.endswith((".xlsx", ".xlsm")):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            sheets = {}
            for ws in wb.worksheets:
                rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
                sheets[ws.title] = [r for r in rows if any(c not in (None, "") for c in r)]
            return {"sheets": sheets}
        if low.endswith(".xls"):
            return {"error": "暂不支持旧版 .xls 文件，请另存为 .xlsx 或 .csv 后上传"}
    except Exception as exc:  # noqa: BLE001
        return {"error": f"文件无法读取（{type(exc).__name__}）：请确认是普通 Excel/CSV 文件"}
    return {"error": "不支持的表格格式"}


# ------------------------------------------------------------------ 识别
def _find_header(rows: list[list]) -> int | None:
    for i, r in enumerate(rows[:10]):
        names = [_split_header(h)[0] for h in r if h]
        hits = sum(1 for n in names if match_field(n) or n in PARAM_NAME_COLS + VALUE_COLS + USE_COLS
                   + sum(MAT_COLS.values(), []) + AMBIGUOUS + SCHEME_COLS)
        if hits >= 2:
            return i
    return None


def _col(header: list[str], names: list[str]) -> int | None:
    for i, h in enumerate(header):
        if _split_header(h)[0] in names:
            return i
    return None


def _role_from_name(name: str | None, index: int) -> str:
    if name and re.search(r"原|基准|调整前|现状|初始", name):
        return "baseline"
    if name and re.search(r"调整|候选|新|比选|对比|方案二|方案2", name):
        return "candidate"
    return "baseline" if index == 0 else "candidate"


def analyze(rows: list[list], *, source: dict) -> dict:
    """生成预览：items（将写入的值）、ambiguous（需选择含义的列）、notes。"""
    preview = {"import_id": new_id("IMP"), "source": source, "layout": None, "items": [], "ambiguous": [],
               "notes": [], "schemes": []}
    hi = _find_header(rows)
    if hi is None:
        preview["notes"].append("没有找到可识别的表头（例如“幕墙宽、总高、分格宽、线重、单价”或“参数名称 | 数值 | 单位”）。"
                                "可以改用对话直接输入，或下载空白模板填写。")
        return preview
    header = [_clean(h) for h in rows[hi]]
    body = rows[hi + 1:]
    sheet = source.get("sheet") or "表格"

    def loc(r, c):
        return f"{source.get('filename') or '粘贴内容'}｜{sheet}｜第 {hi + 2 + r} 行第 {c + 1} 列"

    c_name, c_val, c_unit = _col(header, PARAM_NAME_COLS), _col(header, VALUE_COLS), _col(header, UNIT_COLS)
    c_use = _col(header, USE_COLS)
    c_scheme = _col(header, SCHEME_COLS)
    if c_use is not None and (_col(header, MAT_COLS["linear_mass"]) is not None or _col(header, MAT_COLS["unit_price"])
                              is not None):
        preview["layout"] = "materials"
        cols = {k: _col(header, v) for k, v in MAT_COLS.items()}
        for r, row in enumerate(body):
            use = (row[c_use] if c_use < len(row) else "") or ""
            comp = "column" if "立柱" in use or "柱" == use else "beam" if "横梁" in use or "梁" == use else None
            if comp is None:
                if any(row):
                    preview["notes"].append(f"{loc(r, c_use)}：用途“{use}”不是立柱或横梁，已跳过")
                continue
            role = _role_from_name(row[c_scheme] if c_scheme is not None and c_scheme < len(row) else None, 0)
            tax_txt = row[cols["tax"]] if cols["tax"] is not None and cols["tax"] < len(row) else None
            tax = "inclusive" if tax_txt and "含税" in tax_txt and "不含" not in tax_txt else (
                "exclusive" if tax_txt and re.search(r"未税|不含税", tax_txt) else None)
            for attr in ("name", "linear_mass", "unit_price"):
                ci = cols[attr]
                if ci is None or ci >= len(row) or row[ci] in (None, ""):
                    continue
                _, hunit = _split_header(header[ci])
                unit = hunit
                if attr == "unit_price" and cols["price_unit"] is not None and cols["price_unit"] < len(row):
                    unit = row[cols["price_unit"]] or unit
                if attr == "linear_mass" and not unit:
                    unit = "kg/m" if re.search(r"kg/m", header[ci]) else None
                uf = f"{comp}_{attr}"
                preview["items"].append({"role": role, "field": uf, "raw_value": row[ci], "raw_unit": unit,
                                         "location": loc(r, ci), "tax": tax if attr == "unit_price" else None})
        return _validate(preview)
    if c_name is not None and c_val is not None:
        preview["layout"] = "params"
        names_seen = []
        for r, row in enumerate(body):
            name = row[c_name] if c_name < len(row) else None
            val = row[c_val] if c_val < len(row) else None
            if not name or val in (None, ""):
                continue
            pname, hunit = _split_header(name)
            unit = (row[c_unit] if c_unit is not None and c_unit < len(row) else None) or hunit
            sname = row[c_scheme] if c_scheme is not None and c_scheme < len(row) else None
            if sname not in names_seen:
                names_seen.append(sname)
            role = _role_from_name(sname, names_seen.index(sname))
            f = match_field(pname)
            if f is None:
                preview["ambiguous"].append({"key": f"r{r}", "column": pname, "sample": val, "unit": unit,
                                             "role": role, "location": loc(r, c_val)})
                continue
            preview["items"].append({"role": role, "field": f, "raw_value": val, "raw_unit": unit,
                                     "location": loc(r, c_val), "tax": None})
        preview["schemes"] = [n for n in names_seen if n]
        return _validate(preview)
    preview["layout"] = "wide"
    data_rows = [row for row in body if any(c not in (None, "") for c in row)]
    if len(data_rows) > 2:
        preview["notes"].append(f"表中有 {len(data_rows)} 行方案；当前版本一次只比较两个方案，只导入前两行")
    for r, row in enumerate(data_rows[:2]):
        sname = row[c_scheme] if c_scheme is not None and c_scheme < len(row) else None
        role = _role_from_name(sname, r)
        if sname:
            preview["schemes"].append(sname)
        for ci, h in enumerate(header):
            if ci == c_scheme or ci >= len(row) or row[ci] in (None, ""):
                continue
            name, hunit = _split_header(h)
            f = match_field(name)
            if f is None:
                if name in AMBIGUOUS or name in ("线重", "单价", "价格"):
                    preview["ambiguous"].append({"key": f"r{r}c{ci}", "column": name, "sample": row[ci],
                                                 "unit": hunit, "role": role, "location": loc(r, ci)})
                elif name:
                    preview["notes"].append(f"列“{name}”不是可计算参数，已作为说明忽略")
                continue
            preview["items"].append({"role": role, "field": f, "raw_value": row[ci], "raw_unit": hunit,
                                     "location": loc(r, ci), "tax": None})
    return _validate(preview)


def _validate(preview: dict) -> dict:
    """逐项做单位与数值校验，问题写在条目上，不阻止其他条目导入。"""
    for it in preview["items"]:
        d = USER_FIELDS[it["field"]]
        if d["kind"] == "text":
            it["ok"], it["display"] = True, str(it["raw_value"])
            continue
        v, unit, iss = parse_quantity(it["raw_value"], it["raw_unit"], d["kind"], it["field"])
        if v is not None and v < 0:
            iss = [Issue(Code.VALUE_INVALID, "数值不能为负", it["field"])]
        it["ok"] = not iss
        it["display"] = f"{format(v.normalize(), 'f')} {unit}" if v is not None and not iss else None
        it["problem"] = iss[0].message if iss else None
    found = {it["field"] for it in preview["items"] if it["ok"]}
    preview["missing"] = [USER_FIELDS[f]["label"] for f in ("width_m", "height_m", "grid_width_m", "grid_height_m",
                                                            "column_linear_mass", "beam_linear_mass",
                                                            "column_unit_price", "beam_unit_price") if f not in found]
    return preview


def preview_from_paste(text: str) -> dict | None:
    rows = read_paste(text)
    if not rows:
        return None
    return analyze(rows, source={"kind": "paste", "filename": None, "sheet": "粘贴内容"})


def previews_from_file(filename: str, data: bytes) -> dict:
    """返回 {preview} 或 {sheet_choice: [...]} 或 {error}。"""
    got = read_file(filename, data)
    if "error" in got:
        return got
    cands = []
    for name, rows in got["sheets"].items():
        if rows and _find_header(rows) is not None:
            cands.append((name, rows))
    if not cands:
        return {"error": f"{filename} 中没有找到可识别的表头。可下载空白模板，或直接在对话中输入参数。"}
    if len(cands) > 1:
        return {"sheet_choice": [c[0] for c in cands], "sheets": {n: r for n, r in cands}, "filename": filename}
    name, rows = cands[0]
    return {"preview": analyze(rows, source={"kind": "file", "filename": filename, "sheet": name})}


def template_xlsx() -> bytes:
    """空白模板（不含任何示例数值）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "方案参数"
    ws.append(["参数名称", "数值", "单位", "方案名称"])
    for label, unit in (("幕墙总宽", "m"), ("幕墙总高", "m"), ("分格宽", "m"), ("分格高", "m")):
        ws.append([label, None, unit, None])
    ws2 = wb.create_sheet("材料报价")
    ws2.append(["用途", "名称", "线重(kg/m)", "单价", "价格单位", "出处", "日期", "税费说明"])
    ws2.append(["立柱", None, None, None, "元/kg", None, None, "未税"])
    ws2.append(["横梁", None, None, None, "元/kg", None, None, "未税"])
    ws3 = wb.create_sheet("说明")
    for line in ("只需填写本次需要的内容，缺的项可以之后在对话中补充。",
                 "线重是型材每米的质量（kg/m）；单价按未税材料价，可写 元/kg 或 元/吨。",
                 "“方案名称”可填“原方案”“调整方案”；不填则视为原方案。",
                 "本模板不含任何示例数值，导入前系统会显示识别结果供你确认。"):
        ws3.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
